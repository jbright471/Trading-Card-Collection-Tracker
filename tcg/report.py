"""HTML report generation via Jinja2."""
import json
from datetime import datetime
from pathlib import Path

import openpyxl
from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl.styles import Font, PatternFill

from tcg.config import EXCEL_FILE, HTML_OUTPUT, PLACEHOLDER_IMG

_TEMPLATE_DIR = Path(__file__).parent / 'templates'
_env = Environment(
    loader=FileSystemLoader(str(_TEMPLATE_DIR)),
    autoescape=select_autoescape(['html', 'xml']),
    trim_blocks=True,
    lstrip_blocks=True,
)


def generate_html_report(items, total_value, total_profit_loss,
                          history_data=None, movers_data=None, deals=None,
                          dollar_delta=None, pct_delta=None,
                          gainers=None, losers=None):
    # Group items by game for per-section rendering
    from collections import defaultdict
    by_game = defaultdict(list)
    for item in items:
        by_game[item['data']['game']].append(item)
    game_order = [g for g in ('MTG', 'YGO', 'PKM') if g in by_game]

    template = _env.get_template('report.html')
    rendered = template.render(
        today=datetime.now().strftime('%Y-%m-%d'),
        items=items,
        by_game=by_game,
        game_order=game_order,
        total_value=total_value,
        total_profit_loss=total_profit_loss,
        pl_color="#27ae60" if total_profit_loss >= 0 else "#c0392b",
        pl_sign="+" if total_profit_loss >= 0 else "",
        dollar_delta=dollar_delta,
        pct_delta=pct_delta,
        deals=deals or [],
        gainers=gainers or [],
        losers=losers or [],
        placeholder_img=PLACEHOLDER_IMG,
        history_dates_json=json.dumps([row[0] for row in (history_data or [])]),
        history_values_json=json.dumps([row[1] for row in (history_data or [])]),
        movers_json=json.dumps(movers_data or {}),
    )
    with open(HTML_OUTPUT, 'w', encoding='utf-8') as f:
        f.write(rendered)
    print(f"Report generated: {HTML_OUTPUT}")


def write_excel(items, total_value, total_profit_loss):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Card Prices"
    ws.append(['Date', 'Game', 'Card Name', 'Set', 'Prices', 'Notes'])

    today = datetime.now().strftime('%Y-%m-%d')
    for item in items:
        d = item['data']
        note = f"x{item['quantity']}" if item['quantity'] > 1 else ""
        if item['profit_loss'] is not None:
            sign = "+" if item['profit_loss'] >= 0 else ""
            note += f" (P/L: {sign}${item['profit_loss']:.2f})"
        ws.append([today, d['game'], d['name'], d['set'], item['price_str'], note])

    ws.append(['Total', '', '', '', f"{total_value:.2f}", f"Total P/L: ${total_profit_loss:.2f}"])

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    wb.save(EXCEL_FILE)
    print(f"Saved to {EXCEL_FILE}")
